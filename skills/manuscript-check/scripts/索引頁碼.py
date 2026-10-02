"""索引頁碼：依索引詞表在排版 PDF 找出每個詞出現的書頁，產出索引（Word）與核對用報表（Excel）。只讀不改。

用法：
    python 索引頁碼.py 索引詞表.txt 排版稿.pdf
    python 索引頁碼.py 索引詞表.xlsx 排版稿.pdf --輸出 D:\\索引

也可以把兩個檔案一起拖放到「索引頁碼.bat」上（順序不拘）。
索引詞表格式見 設定\\索引詞表範本.txt；頁碼差、範圍符號等見 設定\\設定.ini 的 [索引]。
"""
import argparse
import os
import sys
import time
from datetime import datetime

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402

from checks import index_pages, pdf_text  # noqa: E402
from checks.common import load_config  # noqa: E402

FONT = '微軟正黑體'


def _sheet(wb, title, header, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(header)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(name=FONT, bold=True)
        c.fill = PatternFill('solid', fgColor='DDE4EE')
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=FONT)
            c.alignment = Alignment(wrap_text=True, vertical='top')
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = 'A2'
    if rows:
        ws.auto_filter.ref = ws.dimensions


def main():
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass
    ap = argparse.ArgumentParser(description='依索引詞表在排版 PDF 抓頁碼（規則式，不使用 AI）')
    ap.add_argument('檔案', nargs=2, help='索引詞表（.txt／.csv／.xlsx）與排版 .pdf（順序不拘）')
    ap.add_argument('--設定', default=os.path.join(HERE, '設定', '設定.ini'))
    ap.add_argument('--輸出', help='輸出資料夾（預設與索引詞表同資料夾）')
    args = ap.parse_args()

    files = [os.path.abspath(f) for f in args.檔案]
    pdfs = [f for f in files if f.lower().endswith('.pdf')]
    lists = [f for f in files if f.lower().endswith(('.txt', '.csv', '.xlsx'))]
    if len(pdfs) != 1 or len(lists) != 1:
        print('需要一個索引詞表（.txt／.csv／.xlsx）和一個 .pdf')
        return 1

    cfg = load_config(args.設定)
    sec = cfg['索引']
    t0 = time.time()
    entries = index_pages.load_terms(lists[0])
    print(f'■ 索引詞表：{os.path.basename(lists[0])}（{len(entries)} 條）')
    pdf = pdf_text.load(pdfs[0], cfg)
    print(f'■ 排版稿：{os.path.basename(pdfs[0])}（{pdf.n_pages} 頁）')
    if not pdf.main.text:
        print('  PDF 沒有文字層（可能是掃描檔），無法抓頁碼')
        return 1
    scan = None
    rng = sec.get('掃描範圍', '').strip()
    if rng:
        a, b = rng.replace('－', '-').split('-')
        scan = (int(a), int(b))
    if sec.get('含註腳', '是') != '是':
        pdf.notes.text, pdf.notes.pages = '', []
    skip_toc = sec.get('略過目次', '是') == '是'
    index_pages.find_pages(entries, pdf, scan, skip_toc)
    toc_note = '、'.join(map(str, sorted(pdf.toc_pages))) if skip_toc and pdf.toc_pages else '無'
    if toc_note != '無':
        print(f'  略過目次頁：PDF 第 {toc_note} 頁')
    lines, table, missing, off, off_note = index_pages.build(entries, pdf, sec)
    print(f'  {off_note}')
    print(f'  找到 {len(entries) - len(missing)} 條，找不到 {len(missing)} 條')

    out_dir = args.輸出 or os.path.dirname(lists[0])
    stem = os.path.splitext(os.path.basename(pdfs[0]))[0]
    stamp = f'{datetime.now():%Y%m%d-%H%M}'
    docx_out = index_pages.write_docx(os.path.join(out_dir, f'{stem}_索引_{stamp}.docx'), lines)

    wb = Workbook()
    ov = wb.active
    ov.title = '總覽'
    for r in [['索引詞表', os.path.basename(lists[0])], ['排版稿', f'{os.path.basename(pdfs[0])}（{pdf.n_pages} 頁）'],
              ['頁碼', off_note], ['掃描範圍', f'PDF 第 {scan[0]}–{scan[1]} 頁' if scan else '全部'],
              ['略過目次頁', toc_note],
              ['結果', f'{len(entries)} 條，找不到 {len(missing)} 條'],
              ['說明', '頁首頁尾不算；英文詞不分大小寫、需完整單字；中文詞沒有詞界，請看「逐筆出處」抽查誤中，'
                     '誤中的用索引詞表的「排除詞」剔除後重跑'],
              ['產生時間', datetime.now().strftime('%Y-%m-%d %H:%M')]]:
        ov.append(r)
    for row in ov.iter_rows():
        for c in row:
            c.font = Font(name=FONT, bold=c.column == 1)
            c.alignment = Alignment(wrap_text=True, vertical='top')
    ov.column_dimensions['A'].width = 12
    ov.column_dimensions['B'].width = 90
    _sheet(wb, '索引', ['索引（可直接複製）'], [[ln] for ln in lines], [80])
    _sheet(wb, '詞條', ['詞條', '搜尋詞', '書頁', '命中次數', '參見', '狀態'], table, [24, 24, 40, 10, 16, 10])
    detail = []
    for e in entries:
        for pno, term, ctx, is_note in sorted(e.hits):
            book = pno - off if off is not None else ''
            detail.append([e.heading, term, pno, book if book != '' and book >= 1 else '（前頁）' if off is not None else '',
                           '註腳' if is_note else '正文', ctx])
    _sheet(wb, '逐筆出處', ['詞條', '搜尋詞', 'PDF 頁', '書頁', '位置', '前後文'], detail, [20, 16, 8, 8, 8, 50])
    xlsx_out = os.path.join(out_dir, f'{stem}_索引核對_{stamp}.xlsx')
    wb.save(xlsx_out)
    print(f'  索引：{docx_out}')
    print(f'  核對報表：{xlsx_out}（{time.time() - t0:.1f} 秒）')
    return 0


if __name__ == '__main__':
    sys.exit(main())
