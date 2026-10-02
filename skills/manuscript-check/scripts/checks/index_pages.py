"""⑥ 索引頁碼：依索引詞表在排版 PDF 裡找出每個詞出現的書頁，排成索引。

- 搜尋範圍：正文與隨頁註腳；頁首頁尾（章名頁眉、頁碼）不算，否則每一頁都會被算進去。
- 書頁：從頁尾／頁首抓到的頁碼推算「PDF 頁 − 書頁」的差，沒印頁碼的頁依此補推；也可在設定裡直接指定。
- 英文詞：不分大小寫、前後不能緊接英文字母（AI 不會對到 MAIN）；行末斷字（gover-／nance）照樣找得到。
- 中文詞沒有詞界，可用「排除詞」剔除誤中（例：找「治理」時排除「管治理論」）。
"""
import os
import re
from collections import Counter
from dataclasses import dataclass, field

from .common import covered
from .pdf_text import normalize


@dataclass
class Entry:
    heading: str                  # 詞條（子條目用「上層>子詞」）
    terms: list                   # 搜尋詞
    see: str = ''                 # 參見
    excludes: list = field(default_factory=list)
    line: int = 0
    pages: set = field(default_factory=set)       # PDF 頁
    note_pages: set = field(default_factory=set)  # 只出現在註腳的 PDF 頁
    hits: list = field(default_factory=list)      # (PDF 頁, 搜尋詞, 前後文, 是否註腳)


def load_terms(path):
    rows = []
    if path.lower().endswith('.xlsx'):
        import openpyxl
        ws = openpyxl.load_workbook(path, read_only=True).worksheets[0]
        data = [[('' if c is None else str(c)).strip() for c in r] for r in ws.iter_rows(values_only=True)]
        head = data[0] if data else []
        col = {name: head.index(name) for name in ('詞條', '搜尋詞', '參見', '排除詞') if name in head}
        if '詞條' not in col:
            raise ValueError('Excel 索引詞表第一列要有「詞條」欄（可另加 搜尋詞、參見、排除詞）')
        for i, r in enumerate(data[1:], 2):
            get = lambda k: r[col[k]] if k in col and col[k] < len(r) else ''  # noqa: E731
            if get('詞條'):
                rows.append((i, get('詞條'), get('搜尋詞'), get('參見'), get('排除詞')))
    else:
        with open(path, encoding='utf-8-sig') as f:
            for i, ln in enumerate(f, 1):
                ln = ln.rstrip('\r\n')
                if not ln.strip() or ln.startswith(("'", '#')):
                    continue
                c = (ln.split(',') + ['', '', ''])[:4]
                rows.append((i, c[0].strip(), c[1].strip(), c[2].strip(), c[3].strip()))
    entries = []
    for i, head, terms, see, exc in rows:
        leaf = head.split('>')[-1].strip()
        given = [t.strip() for t in terms.split('|') if t.strip()]
        # 沒填搜尋詞就找詞條本身；純「參見」條目（有參見、沒搜尋詞）不搜尋
        search = given or ([] if see else [leaf])
        entries.append(Entry(head.strip(), search, see, [e.strip() for e in exc.split('|') if e.strip()], i))
    return entries


def _dehyphen(stream):
    """做一份去掉行末斷字連字號的文字（gover-／nance → governance），並記住每個字對應回原文字流的位置。"""
    out, idx = [], []
    t = stream.text
    for i, ch in enumerate(t):
        if ch == '-' and i in stream.line_ends and i > 0 and t[i - 1].isalpha() and i + 1 < len(t) and t[i + 1].islower():
            continue
        out.append(ch)
        idx.append(i)
    return ''.join(out), idx


def _pattern(term):
    t = re.escape(normalize(term)[0])
    if re.search(r'[A-Za-z]', term):
        return re.compile(r'(?<![A-Za-z])' + t + r'(?![A-Za-z])', re.I)
    return re.compile(t)


def find_pages(entries, pdf, scan=None, skip_toc=True):
    for stream, is_note in ((pdf.main, False), (pdf.notes, True)):
        text, idx = _dehyphen(stream)
        for e in entries:
            for term in e.terms:
                for m in _pattern(term).finditer(text):
                    if e.excludes and covered(text, m.start(), m.end(), e.excludes):
                        continue
                    pno = stream.pages[idx[m.start()]]
                    if scan and not (scan[0] <= pno <= scan[1]):
                        continue
                    if skip_toc and pno in pdf.toc_pages:
                        continue
                    ctx = text[max(0, m.start() - 12):m.start()] + '【' + m.group() + '】' + text[m.end():m.end() + 12]
                    e.hits.append((pno, term, ctx, is_note))
                    if is_note:
                        if pno not in e.pages:
                            e.note_pages.add(pno)
                    else:
                        e.pages.add(pno)
                        e.note_pages.discard(pno)
    return entries


def page_offset(pdf, override=None):
    """回傳 (PDF 頁 − 書頁 的差, 說明)。"""
    if override is not None:
        return override, f'使用設定的頁碼差 {override}（PDF 第 {override + 1} 頁＝書頁 1）'
    diffs = Counter(p - int(v) for p, v in pdf.printed.items() if v.isdigit())
    if not diffs:
        return None, '頁尾頁首沒有抓到頁碼，索引改用 PDF 頁碼；請在設定填「頁碼差」'
    off, n = diffs.most_common(1)[0]
    total = sum(diffs.values())
    note = f'從 {total} 頁的頁碼推得頁碼差 {off}（PDF 第 {off + 1} 頁＝書頁 1）'
    if n < total:
        note += f'；另有 {total - n} 頁的頁碼與此不一致，可能是頁尾有其他數字，請抽查'
    return off, note


def opt(sec, key, default):
    """讀設定值；用引號包起來的值保留前後空格（設定檔會自動吃掉空格）。"""
    v = sec.get(key, default)
    if len(v) >= 2 and v[0] == v[-1] and v[0] in ('"', "'"):
        v = v[1:-1]
    return v


def fmt_pages(pages, note_pages, off, cfg_sec):
    sep = opt(cfg_sec, '頁碼分隔', ', ')
    dash = opt(cfg_sec, '範圍符號', '–')
    note_mark = opt(cfg_sec, '註腳標記', '')
    merge = cfg_sec.get('合併連續頁', '是') == '是'
    items = sorted([(p, False) for p in pages] + [(p, True) for p in note_pages])
    nums = []
    for p, is_note in items:
        book = p - off if off is not None else p
        if off is not None and book < 1:
            continue    # 書頁 1 之前（前頁）不列
        nums.append((book, is_note))
    out, i = [], 0
    while i < len(nums):
        j = i
        if merge:
            while j + 1 < len(nums) and nums[j + 1][0] == nums[j][0] + 1 and not nums[j + 1][1] and not nums[i][1]:
                j += 1
        a = str(nums[i][0]) + (note_mark if nums[i][1] else '')
        out.append(a if j == i else f'{a}{dash}{nums[j][0]}')
        i = j + 1
    return sep.join(out)


STROKE_FILE = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'data', '筆畫表.tsv')
_strokes = None


def strokes():
    """字 -> (總筆畫, 部首序, 部首外筆畫)，資料來自 Unicode Unihan。"""
    global _strokes
    if _strokes is None:
        _strokes = {}
        with open(STROKE_FILE, encoding='utf-8') as f:
            for ln in f:
                if ln.startswith('#'):
                    continue
                c, s, r, x = ln.rstrip('\n').split('\t')
                _strokes[c] = (int(s), int(r), int(x))
    return _strokes


def sort_key(word, latin_last=True):
    """筆畫排序鍵：首字筆畫少的在前；同筆畫依部首、部首外筆畫；首字相同再比第二字，依此類推。
    西文與數字開頭的詞條依字母排（不分大小寫），整組放在中文之後（或之前）。"""
    table = strokes()
    w = word.strip()
    if w and w[0] in table:
        return (1 if latin_last else 2, [table.get(c, (99, 999, 99)) + (ord(c),) for c in w])
    return (2 if latin_last else 1, [(0, 0, 0, ord(c)) for c in w.casefold()])


def sort_entries(entries, latin_last=True):
    """依筆畫排序；子條目排在所屬上層詞條之下，同層彼此排序。"""
    by_parent = {}
    for e in entries:
        parts = [p.strip() for p in e.heading.split('>')]
        by_parent.setdefault('>'.join(parts[:-1]), []).append(e)
    out = []

    def walk(parent):
        kids = sorted(by_parent.get(parent, []), key=lambda e: sort_key(e.heading.split('>')[-1], latin_last))
        for e in kids:
            out.append(e)
            walk('>'.join(p.strip() for p in e.heading.split('>')))
    walk('')
    placed = set(map(id, out))
    out += [e for e in entries if id(e) not in placed]   # 上層詞條不在詞表裡的子條目，照原順序放最後
    return out


def build(entries, pdf, cfg_sec):
    ov = cfg_sec.get('頁碼差', '').strip()
    off, off_note = page_offset(pdf, int(ov) if ov else None)
    gap = opt(cfg_sec, '詞頁分隔', '　')
    order = cfg_sec.get('排序', '原順序')
    latin_last = cfg_sec.get('西文詞條', '放後面') != '放前面'
    stroke_heads = order == '筆畫' and cfg_sec.get('筆畫標題', '否') == '是'
    if order == '筆畫':
        entries = sort_entries(entries, latin_last)
    lines, table, missing = [], [], []
    last_head = None
    leaves = {x.heading.split('>')[-1].strip() for x in entries}
    seen = set()
    for e in entries:
        parts = e.heading.split('>')
        orphan = len(parts) > 1 and '>'.join(parts[:-1]) not in seen
        seen.add(e.heading)
        if stroke_heads and len(parts) == 1 and (e.pages or e.note_pages or e.see
                                                  or any(x.heading.startswith(e.heading + '>') for x in entries)):
            first = parts[0].strip()[:1]
            head = f'{strokes()[first][0]} 畫' if first in strokes() else '西文'
            if head != last_head:
                lines.append(f'【{head}】')
                last_head = head
        indent = '　' * (len(parts) - 1)
        pages = fmt_pages(e.pages, e.note_pages, off, cfg_sec)
        see = ''
        if e.see:
            see = ('；另見 ' if pages else '見 ') + e.see
        has_kids = any(x.heading.startswith(e.heading + '>') for x in entries)
        if not pages and not e.see:
            missing.append(e)
        if pages or see or has_kids:   # 找不到又沒有子條目的詞條不印進索引（「詞條」表會標「找不到」）
            lines.append(f'{indent}{parts[-1].strip()}{gap}{pages}{see}'.rstrip())
        state = '找不到' if not pages and not e.see else ''
        if order == '筆畫' and any(c not in strokes() for c in parts[-1].strip()[:1]) and re.match(r'[^A-Za-z0-9]', parts[-1].strip()):
            state = (state + '；' if state else '') + '首字沒有筆畫資料，排在最後'
        if orphan:
            state = (state + '；' if state else '') + '上層詞條沒有寫在前面'
        if e.see and e.see not in leaves:
            state = (state + '；' if state else '') + f'參見的「{e.see}」不是索引裡的詞條'
        table.append([e.heading, '|'.join(e.terms), pages, len(e.hits), e.see, state])
    return lines, table, missing, off, off_note


def write_docx(path, lines, title='索引'):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt
    d = Document()
    st = d.styles['Normal']
    st.font.name = '新細明體'
    st.element.rPr.rFonts.set(qn('w:eastAsia'), '新細明體')
    st.font.size = Pt(10.5)
    d.add_heading(title, level=1)
    for ln in lines:
        p = d.add_paragraph(ln)
        p.paragraph_format.space_after = Pt(0)
    d.save(path)
    return path
