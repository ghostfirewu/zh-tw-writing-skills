"""把各項檢查結果寫成 Excel 報表：總覽＋每項一張工作表（另附表格者再加一張）。"""
import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

FONT = '微軟正黑體'
HEAD_FILL = PatternFill('solid', fgColor='DDE4EE')
MAX_HITS = 5000   # 每張工作表最多列幾筆，避免異常稿件拖垮 Excel


def _sheet(wb, title, header, rows, widths):
    ws = wb.create_sheet(title[:31])
    ws.append(header)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(name=FONT, bold=True)
        c.fill = HEAD_FILL
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=FONT)
            c.alignment = Alignment(wrap_text=True, vertical='top')
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(ord('A') + i)].width = w
    ws.freeze_panes = 'A2'
    if rows:
        ws.auto_filter.ref = ws.dimensions
    return ws


def write(path, doc, results, extra=()):
    wb = Workbook()
    ov = wb.active
    ov.title = '總覽'
    ov.append(['稿件', os.path.basename(doc.path)])
    for label, value in extra:
        ov.append([label, value])
    ov.append(['產生時間', datetime.now().strftime('%Y-%m-%d %H:%M')])
    ov.append(['說明', '本報表由規則式程式產生，只列候選、不修改稿件；每一筆都要由編輯判斷。'])
    ov.append([])
    ov.append(['項目', '分類', '筆數', '備註'])
    head_row = ov.max_row

    for res in results.values():
        counts = {}
        for h in res.hits:
            counts[h.category] = counts.get(h.category, 0) + 1
        if counts:
            for i, (cat, n) in enumerate(counts.items()):
                ov.append([res.title if i == 0 else '', cat, n, '；'.join(res.notes) if i == 0 else ''])
        else:
            ov.append([res.title, '（無命中）', 0, '；'.join(res.notes)])

        if res.header:
            header, rows = res.header, list(res.rows[:MAX_HITS])
            widths = [8, 14, 18, 12, 14, 50, 50, 40]
        else:
            header = ['分類', '章', '位置', '命中', '前後文', '說明']
            rows = [[h.category, h.para.chapter if h.para else '', h.para.loc if h.para else '全書',
                     h.match, h.context(), h.note] for h in res.hits[:MAX_HITS]]
            widths = [16, 18, 12, 14, 60, 40]
        if len(res.hits) > MAX_HITS:
            rows.append(['（截斷）'] + [''] * (len(header) - 2) + [f'共 {len(res.hits)} 筆，只列前 {MAX_HITS} 筆'])
        _sheet(wb, res.title, header, rows, widths)
        if res.table:
            _sheet(wb, res.table_title or res.title + '附表', res.table[0], res.table[1:],
                   [18] + [12] * (len(res.table[0]) - 2) + [60])

    for row in ov.iter_rows():
        for c in row:
            c.font = Font(name=FONT, bold=(c.row == head_row or (c.column == 1 and c.row < head_row)))
            c.alignment = Alignment(wrap_text=True, vertical='top')
    for c in ov[head_row]:
        c.fill = HEAD_FILL
    for col, w in zip('ABCD', (18, 20, 8, 80)):
        ov.column_dimensions[col].width = w
    wb.save(path)
    return path
