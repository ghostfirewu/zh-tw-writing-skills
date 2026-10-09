"""稿件健檢第一期的驗收測試：對 fixture.docx 跑全部檢查，逐項核對預期命中。

用法：python tests/test_all.py      （全部通過會印 ALL PASS，否則列出失敗項並以代碼 1 結束）
夾具內容見 fixture.json；夾具改了要用 make_fixture.ps1 重建 fixture.docx。
"""
import os
import sys
import tempfile

sys.stdout.reconfigure(encoding="utf-8")

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))

from checks import docx_text, punct, simplified, rare, variants, structure, stats, foreign, refs, report  # noqa: E402
from checks.common import load_config  # noqa: E402

failures = []


def expect(cond, label):
    print(("  ok   " if cond else "  FAIL ") + label)
    if not cond:
        failures.append(label)


def has(hits, category, match=None, part=None):
    for h in hits:
        if h.category == category and (match is None or h.match == match) and (part is None or h.para.part == part):
            return True
    return False


def count(hits, category, match=None):
    return sum(1 for h in hits if h.category == category and (match is None or h.match == match))


cfg = load_config(os.path.join(os.path.dirname(HERE), '設定', '設定.ini'))
cfg['異形詞']['詞表'] = os.path.join(HERE, '測試詞表.txt')
cfg['異形詞']['譯名表'] = ''
doc = docx_text.load(os.path.join(HERE, 'fixture.docx'))

print('[文字層]')
body = '\n'.join(p.text for p in doc.paras)
expect('这里' not in body, '追蹤修訂中被刪除的文字不列入')
expect('插入（未閉合' in body, '追蹤修訂中插入的文字列入')
expect(any(p.part == '註腳' and '以后再議' in p.text for p in doc.paras), '讀得到註腳')
expect([p.level for p in doc.paras if p.level] == [1, 1, 2, 2, 1], '標題層級：1,1,2,2,1')

print('[① 標點]')
r = punct.run(doc, cfg)
expect(has(r.hits, '未閉合', '「'), '未閉合的「')
expect(has(r.hits, '未閉合', '（'), '修訂插入的未閉合（')
expect(count(r.hits, '疊字', '的的') == 1, '「我的的書」報 1 次、「的的確確」不報')
expect(not has(r.hits, '疊字', '確確'), '「的的確確」的「確確」不報')
expect(has(r.hits, '半形標點夾在中文間', ','), '「說,好」的半形逗號')
expect(has(r.hits, '連續標點', '。，'), '「。，」連續標點')

print('[② 簡體與異常字元]')
r = simplified.run(doc, cfg)
expect(len(simplified.gb_only_chars()) == 2380, 'GB2312 有、Big5 無的漢字共 2380 個（實測值）')
expect(count(r.hits, '簡體字', '这') == 1, '「这」只算未刪除的 1 處')
expect(has(r.hits, '簡體字', '说'), '「说」')
expect(count(r.hits, '一簡多繁字組', '以后') == 2, '「以后」2 處（含註腳）')
expect(has(r.hits, '一簡多繁字組', '以后', part='註腳'), '「以后」在註腳的那處標為註腳')
expect(has(r.hits, '錯轉繁', '頭發'), '錯轉繁「頭發」')

fake = docx_text.Doc('x', [docx_text.Para('內文', 1, '酶與多肽；这', 'Normal', None, False, '章')])
r = simplified.run(fake, cfg)
expect([h.match for h in r.hits if h.category == '簡體字'] == ['这'], '台灣正字（酶、肽）不算簡體，这照報')
r = rare.run(fake, cfg)
expect(has(r.hits, '漢字（基本區）', '酶') and '也是簡體字' not in next(h.note for h in r.hits if h.match == '酶'), '罕用字照列「酶」，但不註「也是簡體字」')

print('[排除詞與白名單長詞]')
fake = docx_text.Doc('x', [docx_text.Para('內文', 1, '分成若干部分，干部開會；管理理論與我的的書。', 'Normal', None, False, '章')])
r = simplified.run(fake, cfg)
expect([h.match for h in r.hits if h.category == '一簡多繁字組'] == ['干部'], '「若干部分」排除、單獨的「干部」照報')
r = punct.run(fake, cfg)
expect([h.match for h in r.hits if h.category == '疊字'] == ['的的'], '白名單長詞「管理理論」排除、「的的」照報')

print('[⑩ 罕用字]')
fake = docx_text.Doc('x', [docx_text.Para('內文', 1, '堃啓𠀋的这café—“”', 'Normal', None, False, '章')])
r = rare.run(fake, cfg)
got = {h.match: h for h in r.hits}
expect(set(got) == {'堃', '啓', '𠀋', '这'}, '列出 Big5 以外的字（堃、啓、𠀋、这），不列常用字、重音字母、一般標點')
expect(got['𠀋'].category == '漢字（擴充 B 區）', '𠀋 標為擴充 B 區')
expect('簡體' in got['这'].note, '「这」註明也是簡體字')
expect(r.table[1][0] == '堃' and r.table[1][1] == 'U+5803', '交美編清單附碼位，非簡體的罕用字排前面')

print('[③ 同書異形並存]')
r = variants.run(doc, cfg)
expect(has(r.hits, '並存', '身份／身分'), '身份與身分並存')
expect(not has(r.hits, '並存', '裏／裡'), '裏／裡未出現，不報')

print('[④ 結構]')
r = structure.run(doc, cfg)
expect(has(r.hits, '標題跳號', '第三章'), '第一章之後直接第三章')
expect(has(r.hits, '標題跳號', '3.3'), '3.1 之後直接 3.3')
expect(has(r.hits, '圖表跳號', '圖1-3'), '圖1-1 之後直接圖1-3')
expect(has(r.hits, '引用了但沒有圖說', '圖1-2'), '內文引用圖1-2，但沒有這張圖')
expect(has(r.hits, '圖說未被引用', '圖1-1'), '圖1-1 沒被內文引用')
expect(not has(r.hits, '圖說未被引用', '表1-1'), '表1-1 有被引用，不報')

print('[⑤ 統計]')
r = stats.run(doc, cfg)
expect(count(r.hits, '無標點長串') == 1, '無標點長串 1 處')
expect([row[0] for row in r.table[1:]] == ['第一章　緒論', '第三章　方法', '參考文獻'], '分章統計三章')

print('[⑧ 外文拼寫]')
r = foreign.run(doc, cfg)
expect(has(r.hits, '疑似拼錯', 'Saccado'), 'Saccado 疑似 Saccardo 拼錯')
expect(not has(r.hits, '疑似拼錯', 'Saccardo'), '高頻形 Saccardo 本身不報')

print('[⑨ 參考文獻]')
r = refs.run(doc, cfg)
expect(has(r.hits, '引用了但書目沒有', 'Jones 2018'), 'Jones 2018 有引用、書目沒有')
expect(has(r.hits, '書目有但內文沒引用', 'Lee 2011'), 'Lee 2011 書目有、內文沒引用')
expect(not has(r.hits, '引用了但書目沒有', 'Smith 2010'), 'Smith 2010 對得上')
expect(not has(r.hits, '書目有但內文沒引用', '王小明 2015'), '王小明 2015 對得上')
expect(not any('1990' in h.match for h in r.hits), '「（1990 年代）」不當成引用')

print('[報表]')
out = os.path.join(tempfile.mkdtemp(), 'out.xlsx')
results = {name: mod.run(doc, cfg) for name, mod in [('標點', punct), ('簡體', simplified), ('罕用字', rare), ('異形', variants),
                                                    ('結構', structure), ('統計', stats), ('外文', foreign), ('文獻', refs)]}
report.write(out, doc, results)
import openpyxl  # noqa: E402
wb = openpyxl.load_workbook(out)
expect(wb.sheetnames[0] == '總覽' and len(wb.sheetnames) >= 8, '報表有總覽與各項工作表')

print()
if failures:
    print(f'{len(failures)} FAIL')
    sys.exit(1)
print('ALL PASS')
